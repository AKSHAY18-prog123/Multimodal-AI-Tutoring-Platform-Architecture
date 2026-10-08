import asyncio
import os
from pathlib import Path
import fitz # PyMuPDF
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw, ImageFont

import uuid
from sqlalchemy import select
from backend.app.database.session import AsyncSessionLocal
from backend.app.database.init_db import init_database
from backend.app.database.models.user import User
from backend.app.database.models.learner import LearnerProfile
from backend.app.database.models.course import Course
from backend.app.database.models.document import Document
from backend.app.database.models.knowledge import Topic, Concept, ConceptRelationship
from backend.app.database.models.assessment import Question
from backend.app.workers.document_worker import process_document_background
from backend.app.core.config import settings
from backend.app.core.logging import logger

def generate_sample_pdf(output_path: Path):
    doc = fitz.open()
    page = doc.new_page()

    text = """Chapter 7: Deadlocks & Synchronization

7.1 Deadlock Characterization
In a multiprogramming environment, several processes may compete for a finite number of resources.
A process requests resources; if the resources are not available at that time, the process enters a wait state.
Sometimes, a waiting process is never again able to change state, because the resources it has requested are held by other waiting processes.
This situation is called a deadlock.

Four conditions must hold simultaneously for a deadlock to arise:
1. Mutual Exclusion: At least one resource must be held in a non-shareable mode.
2. Hold and Wait: A process must be currently holding at least one resource and requesting additional resources.
3. No Preemption: Resources cannot be preempted; a resource can be released only voluntarily by the process holding it.
4. Circular Wait: A closed chain of processes exists such that each process holds at least one resource needed by the next.

7.2 Deadlock Avoidance: The Banker's Algorithm
The Banker's Algorithm is a deadlock avoidance algorithm developed by Edsger Dijkstra.
When a new process enters the system, it must declare the maximum number of instances of each resource type that it may need.
This number may not exceed the total number of resources in the system.

Data Structures for Banker's Algorithm:
- Available: Vector of length m. If Available[j] equals k, then k instances of resource type Rj are available.
- Max: n x m matrix. If Max[i][j] equals k, then process Pi may request at most k instances of resource type Rj.
- Allocation: n x m matrix. If Allocation[i][j] equals k, then process Pi is currently allocated k instances of resource type Rj.
- Need: n x m matrix. If Need[i][j] equals k, then process Pi may need k more instances of resource type Rj:
Need[i][j] = Max[i][j] - Allocation[i][j]

Safety Algorithm:
Let Work and Finish be vectors of length m and n, respectively.
1. Initialize Work = Available and Finish[i] = false for all i.
2. Find an index i such that both Finish[i] == false and Need[i] <= Work.
If no such i exists, go to step 4.
3. Work = Work + Allocation[i], Finish[i] = true. Go to step 2.
4. If Finish[i] == true for all i, then the system is in a safe state.
"""
    rect = fitz.Rect(50, 50, 550, 750)
    page.insert_textbox(rect, text, fontsize=11, fontname="helv")
    doc.save(str(output_path))
    doc.close()
    logger.info(f"Generated sample PDF at {output_path}")

def generate_sample_pptx(output_path: Path):
    prs = Presentation()

    # Slide 1: Title
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Module 4: Deadlocks in Operating Systems"
    slide1.placeholders[1].text = "Deadlock Avoidance & The Banker's Algorithm\nComputer Science Department"

    # Slide 2: Banker's Algorithm
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Deadlock Avoidance: Banker's Algorithm"
    body = slide2.placeholders[1]
    tf = body.text_frame
    tf.text = "Key Principles of Banker's Algorithm:"
    p1 = tf.add_paragraph()
    p1.text = "• Developed by Edsger Dijkstra for bank credit limits"
    p2 = tf.add_paragraph()
    p2.text = "• Simulates allocation of maximum possible resource demands"
    p3 = tf.add_paragraph()
    p3.text = "• Validates existence of a Safe Sequence <P1, P2, ... Pn>"
    p4 = tf.add_paragraph()
    p4.text = "• Denies requests that would transition the system into an Unsafe State"

    prs.save(str(output_path))
    logger.info(f"Generated sample PPTX at {output_path}")

def generate_sample_diagram(output_path: Path):
    img = Image.new("RGB", (600, 400), color=(245, 247, 250))
    draw = ImageDraw.Draw(img)

    # Draw boxes for Processes P1, P2 and Resources R1, R2
    # P1 (Process) - Circle
    draw.ellipse([80, 150, 180, 250], fill=(219, 234, 254), outline=(37, 99, 235), width=3)
    draw.text((120, 190), "Process P1", fill=(30, 58, 138))

    # P2 (Process) - Circle
    draw.ellipse([420, 150, 520, 250], fill=(219, 234, 254), outline=(37, 99, 235), width=3)
    draw.text((460, 190), "Process P2", fill=(30, 58, 138))

    # R1 (Resource) - Box
    draw.rectangle([250, 60, 350, 140], fill=(254, 243, 199), outline=(217, 119, 6), width=3)
    draw.text((280, 90), "Resource R1\n(Allocation)", fill=(146, 64, 14))

    # R2 (Resource) - Box
    draw.rectangle([250, 260, 350, 340], fill=(254, 243, 199), outline=(217, 119, 6), width=3)
    draw.text((280, 290), "Resource R2\n(Allocation)", fill=(146, 64, 14))

    # Directed edges (Arrows)
    draw.line([(180, 180), (250, 110)], fill=(220, 38, 38), width=3) # P1 -> R1
    draw.line([(350, 110), (420, 180)], fill=(37, 99, 235), width=3) # R1 -> P2
    draw.line([(420, 220), (350, 290)], fill=(220, 38, 38), width=3) # P2 -> R2
    draw.line([(250, 290), (180, 220)], fill=(37, 99, 235), width=3) # R2 -> P1

    draw.text((150, 20), "Resource Allocation Graph (RAG) with Circular Wait Cycle", fill=(17, 24, 39))

    img.save(str(output_path))
    logger.info(f"Generated sample diagram at {output_path}")

async def seed_sample_course():
    await init_database()

    upload_dir = Path(settings.UPLOAD_DIR)
    img_dir = Path(settings.PROCESSED_DIR) / "images"
    upload_dir.mkdir(parents=True, exist_ok=True)
    img_dir.mkdir(parents=True, exist_ok=True)

    pdf_file = upload_dir / "Operating_Systems_Ch7.pdf"
    pptx_file = upload_dir / "Module_4_Deadlocks.pptx"
    diag_file = img_dir / "rag_deadlock_cycle.png"

    generate_sample_pdf(pdf_file)
    generate_sample_pptx(pptx_file)
    generate_sample_diagram(diag_file)

    async with AsyncSessionLocal() as session:
        # 1. Resolve or Create Demo Student
        u_res = await session.execute(select(User).order_by(User.last_active_at.desc()).limit(1))
        student = u_res.scalar_one_or_none()
        if not student:
            student = User(
                id=str(uuid.uuid4()),
                full_name="Demo Student",
                onboarding_completed=True
            )
            session.add(student)
            profile = LearnerProfile(
                user_id=student.id,
                status="uncalibrated",
                overall_mastery=0.0
            )
            session.add(profile)
            await session.flush()

        # 2. Create Course
        course = Course(
            id="cs301-os",
            user_id=student.id,
            title="CS 301: Modern Operating Systems",
            code="CS301",
            description="Comprehensive university curriculum covering Process Scheduling, Deadlocks, Memory Management, and Concurrency.",
            subject="Computer Science"
        )
        session.add(course)
        await session.flush()

        # 2. Add Documents
        doc_pdf = Document(
            id="doc-os-pdf",
            course_id=course.id,
            filename=pdf_file.name,
            file_path=str(pdf_file),
            file_type="pdf",
            source_category="course_source",
            file_size_bytes=pdf_file.stat().st_size,
            status="uploaded"
        )
        session.add(doc_pdf)

        doc_pptx = Document(
            id="doc-os-pptx",
            course_id=course.id,
            filename=pptx_file.name,
            file_path=str(pptx_file),
            file_type="pptx",
            source_category="course_source",
            file_size_bytes=pptx_file.stat().st_size,
            status="uploaded"
        )
        session.add(doc_pptx)
        await session.commit()

        logger.info("Sample course & documents registered in DB.")

    # Process documents through background ingestion pipeline
    logger.info("Ingesting PDF...")
    await process_document_background("doc-os-pdf")
    logger.info("Ingesting PPTX...")
    await process_document_background("doc-os-pptx")

    logger.info("Sample course ingestion complete!")

if __name__ == "__main__":
    asyncio.run(seed_sample_course())
