import os
from pathlib import Path
from typing import List, Dict, Any
from pptx import Presentation
from backend.app.core.config import settings
from backend.app.core.logging import logger

class PPTIngestionService:
    """Extracts slide text, titles, bullet points, tables, and notes from PPTX while preserving slide numbers."""

    def __init__(self):
        self.image_output_dir = Path(settings.PROCESSED_DIR) / "images"
        self.image_output_dir.mkdir(parents=True, exist_ok=True)

    def extract_pptx(self, file_path: str, document_id: str) -> List[Dict[str, Any]]:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"PPTX file not found: {file_path}")

        prs = Presentation(file_path)
        slides_data = []

        for slide_index, slide in enumerate(prs.slides):
            slide_number = slide_index + 1
            title = None
            bullet_points = []
            raw_text_parts = []
            tables_data = []
            images = []
            notes = ""

            # Extract slide notes if available
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()

            # Iterate shapes
            for shape_idx, shape in enumerate(slide.shapes):
                # Title
                if shape == slide.shapes.title and shape.has_text_frame:
                    title = shape.text_frame.text.strip()
                    raw_text_parts.append(title)
                    continue

                # Text frames
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if not text:
                            continue
                        raw_text_parts.append(text)
                        if paragraph.level > 0 or text.startswith(("•", "-", "*")):
                            bullet_points.append(text)

                # Tables
                if shape.has_table:
                    table = shape.table
                    rows = []
                    for row in table.rows:
                        row_vals = [cell.text.strip() for cell in row.cells]
                        rows.append(row_vals)
                    tables_data.append(rows)

                # Embedded images in slide
                if shape.shape_type == 13: # MSO_SHAPE_TYPE.PICTURE
                    try:
                        image = shape.image
                        img_bytes = image.blob
                        img_ext = image.ext
                        img_filename = f"{document_id}_s{slide_number}_img{shape_idx + 1}.{img_ext}"
                        img_path = self.image_output_dir / img_filename

                        with open(img_path, "wb") as f_img:
                            f_img.write(img_bytes)

                        images.append({
                            "image_path": str(img_path),
                            "slide_number": slide_number,
                            "format": img_ext
                        })
                    except Exception as e:
                        logger.debug(f"Failed to extract image from slide {slide_number}: {e}")

            slides_data.append({
                "slide_number": slide_number,
                "title": title or f"Slide {slide_number}",
                "bullet_points": bullet_points,
                "raw_text": "\n".join(raw_text_parts),
                "notes": notes,
                "tables": tables_data,
                "shape_count": len(slide.shapes),
                "images": images
            })

        logger.info(f"Extracted {len(slides_data)} slides from PPTX {path.name}")
        return slides_data

ppt_ingestor = PPTIngestionService()
